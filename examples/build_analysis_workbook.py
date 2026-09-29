"""Build the analysis workbook: method, data, variables, every trade, results.

    python examples/build_analysis_workbook.py --data <1-minute csv> --out STDVBOT_Analysis.xlsx

Laid out like an experiment (materials, controlled / independent / dependent
variables, procedure) so the analysis can be presented and checked. Trade
rows come from running each strategy; everything derived from them is a live
Excel formula, so changing a controlled variable (e.g. commission) updates
every result.

The workbook embeds the full candle dataset -- keep it out of the repo
(Databento data license).
"""
import argparse
import datetime as dt
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from compare_stdv_strategies import load, mff_rolling  # noqa: E402
from stdvbot import ipda_strategy, killzone_po3_strategy, po3_strategy  # noqa: E402
from stdvbot.backtest import run_backtest  # noqa: E402
from stdvbot.legs import MNQ_TICK_SIZE, MNQ_TICK_VALUE  # noqa: E402
from stdvbot.propfirm import MFF_PRO_50K  # noqa: E402
from stdvbot.strategies import get_strategy  # noqa: E402

FONT = "Arial"
BLUE, GREEN, BLACK = "0000FF", "008000", "000000"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
USD = '$#,##0.00;($#,##0.00);"-"'
USD0 = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
DT = "yyyy-mm-dd hh:mm"
DATE = "yyyy-mm-dd"
PRICE = "#,##0.00"

COMMISSION = 1.24
CONTRACTS = 1


def font(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=FONT, color=color, bold=bold, size=size, italic=italic)


def put(ws, ref, value, color=BLACK, bold=False, fmt=None, fill=None, size=10, wrap=False, italic=False):
    c = ws[ref]
    c.value = value
    c.font = font(color, bold, size, italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return c


def header_row(ws, row, headers, col=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=col + i, value=h)
        c.font = font("FFFFFF", bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")


def section(ws, ref, text, span=6):
    put(ws, ref, text, bold=True, size=11, fill=SECTION_FILL)
    r = ws[ref].row
    for c in range(ws[ref].column + 1, ws[ref].column + span):
        ws.cell(row=r, column=c).fill = SECTION_FILL


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# ---------------------------------------------------------------- strategies
STRATEGIES = [
    ("v1 Manipulation Leg", "Killzone manipulation leg (09:30 NY / 20:00 Asia), fade at inverse-Fib levels",
     "First 6 min after each killzone open; leg of 3-5 one-minute candles, or 1-2 candles at least 3x normal size",
     "Price touches a projected level (1, 2, 2.25, 2.5, 4, 4.5 x leg); levels below 4.5 also need a candlestick pattern",
     "Next Fib level out", "Session VWAP, else 2x risk", "Leg must run against the daily bias; trending days: 4.5 only",
     "Trader's own method (docs/manipulation_leg_strategy.md)"),
    ("v2 Manipulation Leg", "Same as v1, plus a fallback on days with no daily bias",
     "Same as v1", "Same as v1; on no-bias days only the 4.5 level", "Same as v1", "Same as v1",
     "No-bias days only occur during the first 20 days of data (warm-up)", "Fork of v1"),
    ("PO3 Silver Bullet Zone", "4-hour Power of Three: fake move against the candle open, then the real move",
     "Each 4H candle (02/06/10/14/18/22); price trades against the open, then a 5-min close breaks the last swing (MSS)",
     "Retrace into a price gap (FVG) inside the 1-1.5 STDV zone of the manipulation's last leg",
     "Beyond the manipulation extreme", "2.5 STDV (extends to 4 on a strong close)",
     "One trade per 4H candle", "'Standard Deviation + Power of Three' PDF"),
    ("PO3 any gap after zone", "Same as PO3, looser entry",
     "Same as PO3", "Retrace into any price gap from the move, once price has reached the 1-1.5 zone",
     "Same as PO3", "Same as PO3", "The PDF's first entry scenario", "'Standard Deviation + Power of Three' PDF"),
    ("IPDA 12-hour Range", "Prior 12-hour range: fade overextension, or continue from a discount/premium gap",
     "Every 4H boundary: range high/low of last 12h, midpoint, order-flow direction, STDV off the leg into the extreme",
     "Reversal: close back inside 2.0 (or inside 4 after a close beyond 2.5). Continuation: retrace into a gap on the discount/premium side",
     "Reversal: beyond the extreme reached. Continuation: beyond the range extreme",
     "Reversal: range midpoint. Continuation: far side of the range", "Setups live 4 hours; one trade per boundary",
     "'Liquidity Profiles + Standard Deviation Theory' PDF"),
    ("Killzone legs + PO3 entry", "v1's killzone setups, entered with PO3's trigger",
     "Same as v1; price must reach the leg's 2.0 level (4.5 on trending days) within 60 min",
     "Market structure shift, then retrace into a gap inside the 1-1.5 STDV zone", "Beyond the manipulation extreme",
     "2.5 STDV (extends to 4)", "Entry within 4 hours of the killzone", "Hybrid of both"),
]

EXIT_NAMES = {"stop": "Stop", "target": "Target", "target_4": "Target 4 (extended)", "timeout": "Time limit",
              None: "Open at end of data"}
IPDA_NAMES = {"continuation": "Continuation", "reversal_2_2.5": "Reversal at 2-2.5",
              "reversal_terminus": "Reversal at 4 (terminus)"}


def killzone_of(ts):
    return "NY killzone (09:30)" if 9 <= ts.hour < 17 else "Asia killzone (20:00)"


def run_all(df):
    runs = [
        (STRATEGIES[0][0], get_strategy("manipulation_leg").generate_signals(df), None),
        (STRATEGIES[1][0], get_strategy("manipulation_leg_v2").generate_signals(df), None),
    ]
    pos, log = po3_strategy.run(df)
    runs.append((STRATEGIES[2][0], pos, log))
    pos, log = po3_strategy.run(df, entry_mode="any_irl")
    runs.append((STRATEGIES[3][0], pos, log))
    pos, log = ipda_strategy.run(df)
    runs.append((STRATEGIES[4][0], pos, log))
    funnel = {}
    pos, log = killzone_po3_strategy.run(df, funnel=funnel)
    runs.append((STRATEGIES[5][0], pos, log))

    gaps = df.index[(df.index.to_series().diff() > pd.Timedelta(minutes=60)).to_numpy()]
    rows = []
    for name, signals, log in runs:
        trades = run_backtest(df, signals, fee_bps=0.0, slippage_bps=0.0).trades
        if log is not None and len(log) != len(trades):
            raise RuntimeError(f"{name}: strategy log ({len(log)}) and backtester trades ({len(trades)}) differ")
        for i, t in enumerate(trades.itertuples()):
            entry, exit_ = pd.Timestamp(t.entry_time), pd.Timestamp(t.exit_time)
            rec = log[i] if log is not None else None
            if rec is None:
                setup, stop, target, reason = killzone_of(entry), "not logged", "not logged", "not logged"
            else:
                if "profile" in rec:
                    setup = IPDA_NAMES[rec["profile"]]
                elif "killzone" in rec:
                    setup = "NY killzone (09:30)" if rec["killzone"] == "ny" else "Asia killzone (20:00)"
                else:
                    setup = f"4H candle {pd.Timestamp(rec['candle']):%H:00}"
                stop, target, reason = rec["stop"], rec["target"], EXIT_NAMES[rec.get("exit_reason")]
            rows.append({
                "strategy": name, "setup": setup, "direction": t.side.capitalize(),
                "entry_time": entry.to_pydatetime(), "entry_price": float(t.entry_price),
                "stop": stop, "target": target, "exit_time": exit_.to_pydatetime(),
                "exit_price": float(t.exit_price), "exit_reason": reason,
                "entry_date": entry.date(), "exit_date": exit_.date(),
                "held_gap": "Yes" if ((gaps > entry) & (gaps <= exit_)).any() else "No",
            })
    trades = pd.DataFrame(rows)
    point_value = MNQ_TICK_VALUE / MNQ_TICK_SIZE
    sign = trades["direction"].map({"Long": 1, "Short": -1})
    trades["net"] = (trades["exit_price"] - trades["entry_price"]) * sign * point_value * CONTRACTS - COMMISSION * CONTRACTS
    return trades, funnel


# ------------------------------------------------------------------- builder
def build(df, trades, funnel, out_path, source_name):
    wb = Workbook()
    wb._named_styles["Normal"].font = Font(name=FONT, size=10)
    names = [s[0] for s in STRATEGIES]
    trading_days = pd.DatetimeIndex(df.index.normalize().unique())

    ws_over = wb.active
    ws_over.title = "Overview"
    ws_meth = wb.create_sheet("Method")
    ws_data = wb.create_sheet("Data")
    ws_cv = wb.create_sheet("Controlled Variables")
    ws_str = wb.create_sheet("Strategies")
    ws_tl = wb.create_sheet("Trade Log")
    ws_res = wb.create_sheet("Results")
    ws_mon = wb.create_sheet("Monthly")
    ws_eq = wb.create_sheet("Daily Equity")
    ws_brk = wb.create_sheet("Breakdowns")
    ws_fun = wb.create_sheet("Funnel")

    # ---- Data --------------------------------------------------------------
    DATA_FIRST = 9
    data_last = DATA_FIRST + len(df) - 1
    put(ws_data, "A1", "Data (materials): 1-minute price candles", bold=True, size=14)
    meta = [
        ("Source", "Databento, CME Globex MDP 3.0 (Nasdaq-100 futures), OHLCV-1m schema, parent symbol NQ.FUT"),
        ("File", source_name),
        ("Cleaning", "Removed calendar-spread contracts; kept the most-traded contract each day; back-adjusted at contract "
                     "rolls so the switch doesn't create a fake price jump; times converted to New York (UTC-4)."),
        ("Prices", "NQ (E-mini) price series. MNQ (Micro) trades at the same price with $2/point instead of $20, "
                   "so results are calculated in MNQ dollars."),
        ("Rows", f"One row per minute: {len(df):,} candles (count checked on 'Controlled Variables')."),
    ]
    for i, (k, v) in enumerate(meta, start=2):
        put(ws_data, f"A{i}", k, bold=True)
        put(ws_data, f"B{i}", v)
    header_row(ws_data, DATA_FIRST - 1, ["Date / Time (New York)", "Open", "High", "Low", "Close", "Volume"])
    blue = font(BLUE)
    for r, (ts, o, h, lo, c, v) in enumerate(
        zip(df.index.to_pydatetime(), df["open"], df["high"], df["low"], df["close"], df["volume"]), start=DATA_FIRST
    ):
        cells = ((1, ts, DT), (2, float(o), PRICE), (3, float(h), PRICE), (4, float(lo), PRICE),
                 (5, float(c), PRICE), (6, float(v), "#,##0"))
        for col, val, fmt in cells:
            cell = ws_data.cell(row=r, column=col, value=val)
            cell.font = blue
            cell.number_format = fmt
    ws_data.freeze_panes = f"A{DATA_FIRST}"
    widths(ws_data, {"A": 20, "B": 12, "C": 12, "D": 12, "E": 12, "F": 10})

    # ---- Controlled variables ---------------------------------------------
    put(ws_cv, "A1", "Controlled Variables: held identical for every strategy", bold=True, size=14)
    put(ws_cv, "A2", "Blue = input you can change (yellow = most useful to vary). Black = calculated.", italic=True)
    header_row(ws_cv, 4, ["Setting", "Value", "Notes / source"])
    cv_rows = [
        ("point_value", "Dollars per point (MNQ)", MNQ_TICK_VALUE / MNQ_TICK_SIZE, "$#,##0.00",
         "CME contract spec: $0.50 per 0.25-point tick", False),
        ("tick", "Tick size (points)", MNQ_TICK_SIZE, "0.00", "CME contract spec", False),
        ("contracts", "Contracts per trade", CONTRACTS, "0", "Fixed size so strategies compare fairly", True),
        ("commission", "Commission per round trip ($)", COMMISSION, "$#,##0.00",
         "Assumption: typical MNQ round-turn cost; broker rates vary ($0.50-$1.74)", True),
        ("start_bal", "MFF starting balance ($)", MFF_PRO_50K.starting_balance, USD0, "MyFundedFutures Pro 50K", False),
        ("mll", "MFF max loss limit ($, trailing)", MFF_PRO_50K.max_loss_limit, USD0,
         "From the account holder; trails end-of-day highs", False),
        ("lock", "MFF drawdown lock point ($)", MFF_PRO_50K.drawdown_lock_point, USD0,
         "From the account holder: the trailing limit stops rising here", False),
        ("target", "MFF profit target ($)", MFF_PRO_50K.profit_target, USD0, "From the account holder", False),
        ("consistency", "MFF consistency rule (max best-day share)", MFF_PRO_50K.consistency_rule_pct, "0%",
         "From the account holder: best day must be at most 50% of total profit", False),
    ]
    CV = {}
    r = 5
    for key, label, value, fmt, note, highlight in cv_rows:
        put(ws_cv, f"A{r}", label)
        put(ws_cv, f"B{r}", value, color=BLUE, fmt=fmt, fill=INPUT_FILL if highlight else None)
        put(ws_cv, f"C{r}", note)
        CV[key] = f"'Controlled Variables'!$B${r}"
        r += 1
    data_rng = f"Data!$A${DATA_FIRST}:$A${data_last}"
    calc_rows = [
        ("first", "First candle", f"=MIN({data_rng})", DT, "From the Data tab"),
        ("last", "Last candle", f"=MAX({data_rng})", DT, "From the Data tab"),
        ("candles", "Number of candles", f"=COUNT({data_rng})", "#,##0", "From the Data tab"),
        ("weeks", "Weeks tested", "=({last}-{first})/7", "0.0", "Calendar weeks between first and last candle"),
        ("mid", "Halfway point (for stability check)", "={first}+({last}-{first})/2", DT,
         "Trades entered before this = 1st half"),
    ]
    for key, label, formula, fmt, note in calc_rows:
        put(ws_cv, f"A{r}", label)
        f = formula.format(**{k: v for k, v in CV.items()})
        put(ws_cv, f"B{r}", f, fmt=fmt)
        put(ws_cv, f"C{r}", note)
        CV[key] = f"'Controlled Variables'!$B${r}"
        r += 1
    r += 1
    section(ws_cv, f"A{r}", "Rules applied identically to every strategy", span=3)
    for text in [
        "Same data period and candles for every strategy.",
        "A signal is decided on a candle's close and filled at the next candle (no trading on information not yet available).",
        "Exits are filled at candle closes the same way (no special stop/target fills) -- fair between strategies, optimistic vs. live trading.",
        "One position at a time; always at least one flat candle between trades.",
        "Profit is booked on the day a trade exits.",
    ]:
        r += 1
        put(ws_cv, f"A{r}", "- " + text)
    widths(ws_cv, {"A": 44, "B": 18, "C": 80})

    # ---- Strategies -------------------------------------------------------
    put(ws_str, "A1", "Independent Variable: the strategy (the only thing that changes between tests)", bold=True, size=14)
    header_row(ws_str, 3, ["Strategy", "Idea", "When it looks for a setup", "Entry trigger", "Stop", "Target",
                           "Filters / limits", "Source"])
    for i, row in enumerate(STRATEGIES, start=4):
        for j, val in enumerate(row, start=1):
            c = ws_str.cell(row=i, column=j, value=val)
            c.font = font(BLUE if j == 1 else BLACK, bold=(j == 1))
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws_str.row_dimensions[i].height = 75
    r = 4 + len(STRATEGIES) + 1
    section(ws_str, f"A{r}", "Shared math: the STDV / inverse-Fibonacci projection", span=8)
    for text in [
        "Every strategy projects price levels from a price leg: level = leg start + m x (leg start - leg end), for m = 1, 1.5, 2, 2.25, 2.5, 4, 4.5.",
        "The levels land past the leg's starting point, on the opposite side from where the leg went.",
        "2-2.5 = reversal zone; 1-1.5 = 'Silver Bullet Zone' (re-entry area); 4 = 'terminus'; 4.5 = the trader's A+ level.",
        "The two PDFs' 'Standard Deviation' tool and the trader's inverse Fibonacci are the same formula (verified by an automated test).",
    ]:
        r += 1
        put(ws_str, f"A{r}", "- " + text)
    widths(ws_str, {"A": 24, "B": 34, "C": 40, "D": 40, "E": 22, "F": 24, "G": 30, "H": 26})

    # ---- Trade log ----------------------------------------------------------
    TL_FIRST = 5
    tl_last = TL_FIRST + len(trades) - 1
    put(ws_tl, "A1", "Trade Log: every trade from every strategy (raw results)", bold=True, size=14)
    put(ws_tl, "A2", "Blue columns = recorded by the backtest. Black columns = calculated here from the Controlled "
                     "Variables. Times are New York, candle start times.", italic=True)
    tl_headers = ["Strategy", "Setup", "Direction", "Entry time", "Entry price", "Stop", "Target", "Exit time",
                  "Exit price", "Exit reason", "Entry date", "Exit date", "Points", "Gross $", "Commission $",
                  "Net $", "Hold (min)", "Win (1/0)", "Held through daily break / weekend?",
                  "Running net $ (this strategy)", "Running peak $", "Drawdown $"]
    header_row(ws_tl, TL_FIRST - 1, tl_headers)
    ws_tl.row_dimensions[TL_FIRST - 1].height = 45
    blocks = {}
    for i, t in enumerate(trades.itertuples(), start=TL_FIRST):
        blocks.setdefault(t.strategy, [i, i])[1] = i
    for i, t in enumerate(trades.itertuples(), start=TL_FIRST):
        start = blocks[t.strategy][0]
        values = [
            ("A", t.strategy, None), ("B", t.setup, None), ("C", t.direction, None),
            ("D", t.entry_time, DT), ("E", t.entry_price, PRICE),
            ("F", t.stop, PRICE if not isinstance(t.stop, str) else None),
            ("G", t.target, PRICE if not isinstance(t.target, str) else None),
            ("H", t.exit_time, DT), ("I", t.exit_price, PRICE), ("J", t.exit_reason, None),
            ("K", t.entry_date, DATE), ("L", t.exit_date, DATE), ("S", t.held_gap, None),
        ]
        for col, val, fmt in values:
            put(ws_tl, f"{col}{i}", val, color=BLUE, fmt=fmt)
        formulas_ = [
            ("M", f'=IF(C{i}="Long",I{i}-E{i},E{i}-I{i})', "#,##0.00"),
            ("N", f"=M{i}*{CV['point_value']}*{CV['contracts']}", USD),
            ("O", f"={CV['commission']}*{CV['contracts']}", USD),
            ("P", f"=N{i}-O{i}", USD),
            ("Q", f"=ROUND((H{i}-D{i})*1440,0)", "#,##0"),
            ("R", f"=IF(P{i}>0,1,0)", "0"),
            ("T", f"=SUM($P${start}:P{i})", USD),
            ("U", f"=MAX(0,MAX($T${start}:T{i}))", USD),
            ("V", f"=U{i}-T{i}", USD),
        ]
        for col, f, fmt in formulas_:
            put(ws_tl, f"{col}{i}", f, fmt=fmt)
    ws_tl.freeze_panes = f"B{TL_FIRST}"
    widths(ws_tl, {"A": 24, "B": 22, "C": 9, "D": 16, "E": 11, "F": 11, "G": 11, "H": 16, "I": 11, "J": 18,
                   "K": 11, "L": 11, "M": 9, "N": 11, "O": 11, "P": 11, "Q": 9, "R": 8, "S": 14, "T": 13,
                   "U": 12, "V": 11})

    def TL(col):
        return f"'Trade Log'!${col}${TL_FIRST}:${col}${tl_last}"

    # ---- Daily equity -------------------------------------------------------
    EQ_FIRST = 5
    eq_last = EQ_FIRST + len(trading_days) - 1
    put(ws_eq, "A1", "Daily Equity: profit booked each trading day, and running total (1 MNQ contract)", bold=True, size=14)
    put(ws_eq, "A2", "Trading days include Sunday evening sessions (futures reopen Sunday 18:00 New York).", italic=True)
    header_row(ws_eq, EQ_FIRST - 1, ["Date"] + [h for n in names for h in (f"{n}: daily $", f"{n}: running $")])
    ws_eq.row_dimensions[EQ_FIRST - 1].height = 60
    EQ_COLS = {}
    for s_i, n in enumerate(names):
        EQ_COLS[n] = (get_column_letter(2 + 2 * s_i), get_column_letter(3 + 2 * s_i))
    for r_i, day in enumerate(trading_days, start=EQ_FIRST):
        put(ws_eq, f"A{r_i}", day.date(), color=BLUE, fmt=DATE)
        for n in names:
            dcol, ccol = EQ_COLS[n]
            name_ref = f"'Results'!$A${6 + names.index(n)}"
            put(ws_eq, f"{dcol}{r_i}", f"=SUMIFS({TL('P')},{TL('A')},{name_ref},{TL('L')},$A{r_i})", fmt=USD)
            prev = f"{ccol}{r_i - 1}+" if r_i > EQ_FIRST else ""
            put(ws_eq, f"{ccol}{r_i}", f"={prev}{dcol}{r_i}", fmt=USD)
    ws_eq.freeze_panes = f"B{EQ_FIRST}"
    widths(ws_eq, {"A": 12, **{get_column_letter(c): 13 for c in range(2, 2 + 2 * len(names))}})
    chart = LineChart()
    chart.title = "Running net profit by strategy (1 MNQ contract)"
    chart.y_axis.title = "Net $"
    chart.x_axis.title = "Date"
    chart.x_axis.number_format = "mmm d"
    chart.height, chart.width = 11, 26
    for n in names:
        ccol = EQ_COLS[n][1]
        col_idx = ws_eq[f"{ccol}{EQ_FIRST}"].column
        chart.add_data(Reference(ws_eq, min_col=col_idx, min_row=EQ_FIRST - 1, max_row=eq_last), titles_from_data=True)
    chart.set_categories(Reference(ws_eq, min_col=1, min_row=EQ_FIRST, max_row=eq_last))
    ws_eq.add_chart(chart, f"{get_column_letter(3 + 2 * len(names))}4")

    # ---- Results ------------------------------------------------------------
    put(ws_res, "A1", "Results: Dependent Variables (what we measured)", bold=True, size=14)
    put(ws_res, "A2", "Every number is a formula over the Trade Log, except the blue MFF-evaluation counts "
                      "(computed by the evaluation replay described below).", italic=True)
    section(ws_res, "A4", "Performance (1 MNQ contract, after commission)", span=14)
    res_headers = ["Strategy", "Trades", "Trades per week", "Wins", "Win rate", "Winning trades $",
                   "Losing trades $", "Profit factor", "Net profit $", "Avg per trade $", "Avg win $", "Avg loss $",
                   "Max drawdown $", "Drawdown vs MFF limit"]
    header_row(ws_res, 5, res_headers)
    ws_res.row_dimensions[5].height = 45
    R0 = 6
    for i, n in enumerate(names):
        r = R0 + i
        put(ws_res, f"A{r}", n, color=BLUE, bold=True)
        s = f"$A{r}"
        cells = [
            ("B", f"=COUNTIFS({TL('A')},{s})", "0"),
            ("C", f"=IFERROR(B{r}/{CV['weeks']},0)", "0.0"),
            ("D", f'=COUNTIFS({TL("A")},{s},{TL("P")},">0")', "0"),
            ("E", f"=IFERROR(D{r}/B{r},0)", PCT),
            ("F", f'=SUMIFS({TL("P")},{TL("A")},{s},{TL("P")},">0")', USD),
            ("G", f'=-SUMIFS({TL("P")},{TL("A")},{s},{TL("P")},"<=0")', USD),
            ("H", f'=IFERROR(F{r}/G{r},"n/a")', "0.00"),
            ("I", f"=SUMIFS({TL('P')},{TL('A')},{s})", USD),
            ("J", f"=IFERROR(I{r}/B{r},0)", USD),
            ("K", f'=IFERROR(AVERAGEIFS({TL("P")},{TL("A")},{s},{TL("P")},">0"),0)', USD),
            ("L", f'=IFERROR(AVERAGEIFS({TL("P")},{TL("A")},{s},{TL("P")},"<=0"),0)', USD),
            ("M", f"=_xlfn.MAXIFS({TL('V')},{TL('A')},{s})", USD),
            ("N", f'=IF(M{r}>{CV["mll"]},"Exceeds $2,000 limit","Within limit")', None),
        ]
        for col, f, fmt in cells:
            put(ws_res, f"{col}{r}", f, fmt=fmt)

    R1 = R0 + len(names) + 2
    section(ws_res, f"A{R1}", "Stability & robustness: is the result steady, or a few lucky trades?", span=14)
    header_row(ws_res, R1 + 1, ["Strategy", "1st half net $", "2nd half net $", "Median trade $",
                                "Net without top 5 winners $", "Net from trades held through break/weekend $",
                                "Net from all other trades $"])
    ws_res.row_dimensions[R1 + 1].height = 60
    for i, n in enumerate(names):
        r = R1 + 2 + i
        put(ws_res, f"A{r}", f"=A{R0 + i}", color=BLACK, bold=True)
        s = f"$A{r}"
        b = blocks.get(n)
        blk = f"'Trade Log'!$P${b[0]}:$P${b[1]}" if b else None
        top5 = "-".join(f"LARGE({blk},{k})" for k in range(1, 6)) if b else None
        cells = [
            ("B", f'=SUMIFS({TL("P")},{TL("A")},{s},{TL("D")},"<"&{CV["mid"]})', USD),
            ("C", f'=SUMIFS({TL("P")},{TL("A")},{s},{TL("D")},">="&{CV["mid"]})', USD),
            ("D", f"=MEDIAN({blk})" if b else '="n/a"', USD),
            ("E", f'=IF(B{R0 + i}>=5,I{R0 + i}-{top5},"n/a (fewer than 5 trades)")' if b else '="n/a"', USD),
            ("F", f'=SUMIFS({TL("P")},{TL("A")},{s},{TL("S")},"Yes")', USD),
            ("G", f'=SUMIFS({TL("P")},{TL("A")},{s},{TL("S")},"No")', USD),
        ]
        for col, f, fmt in cells:
            put(ws_res, f"{col}{r}", f, fmt=fmt)

    R2 = R1 + 2 + len(names) + 1
    section(ws_res, f"A{R2}", "MyFundedFutures Pro 50K evaluation", span=14)
    header_row(ws_res, R2 + 1, ["Strategy", "Best single day $", "Best day / total profit", "Consistency rule (<= 50%)",
                                "Evaluations started", "Passed", "Failed (hit max loss)", "Unresolved at end of data",
                                "Pass rate (of resolved)", "Median trading days to pass"])
    ws_res.row_dimensions[R2 + 1].height = 60
    mff_rows = {}
    for i, n in enumerate(names):
        r = R2 + 2 + i
        mff_rows[n] = r
        put(ws_res, f"A{r}", f"=A{R0 + i}", bold=True)
        dcol = EQ_COLS[n][0]
        put(ws_res, f"B{r}", f"=MAX('Daily Equity'!${dcol}${EQ_FIRST}:${dcol}${eq_last})", fmt=USD)
        put(ws_res, f"C{r}", f'=IF(I{R0 + i}>0,B{r}/I{R0 + i},"n/a (no profit)")', fmt=PCT)
        put(ws_res, f"D{r}", f'=IF(I{R0 + i}<=0,"n/a",IF(C{r}<={CV["consistency"]},"OK","Fails"))')
        t = trades[trades["strategy"] == n][["exit_time", "net"]]
        roll = mff_rolling(t, trading_days) if len(t) else {"starts": len(trading_days), "passed": 0, "failed": 0,
                                                             "unresolved": len(trading_days), "median_days_to_pass": None}
        put(ws_res, f"E{r}", roll["starts"], color=BLUE, fmt="0")
        put(ws_res, f"F{r}", roll["passed"], color=BLUE, fmt="0")
        put(ws_res, f"G{r}", roll["failed"], color=BLUE, fmt="0")
        put(ws_res, f"H{r}", roll.get("unresolved", 0), color=BLUE, fmt="0")
        put(ws_res, f"I{r}", f'=IFERROR(F{r}/(F{r}+G{r}),"n/a (none resolved)")', fmt=PCT)
        put(ws_res, f"J{r}", roll["median_days_to_pass"] if roll["median_days_to_pass"] is not None else "n/a",
            color=BLUE)
    note_r = R2 + 2 + len(names)
    put(ws_res, f"A{note_r}", "How the evaluation counts were made: for every trading day in the data, a fresh $50,000 "
                              "evaluation starts that day and each strategy's daily profit is replayed forward under the "
                              "rules on 'Controlled Variables' (trailing $2,000 max loss locking at $52,100, $3,000 target, "
                              "50% consistency). 'Unresolved' = neither passed nor failed before the data ran out. "
                              "Computed by examples/compare_stdv_strategies.py (mff_rolling).", italic=True)
    ws_res.merge_cells(f"A{note_r}:N{note_r}")
    ws_res[f"A{note_r}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws_res.row_dimensions[note_r].height = 48
    ws_res[f"E{R2 + 1}"].comment = Comment("Hardcoded: from the evaluation replay (see note below the table).", "STDVBOT")
    widths(ws_res, {"A": 26, **{get_column_letter(c): 14 for c in range(2, 15)}})
    ws_res.freeze_panes = "B6"

    # ---- Monthly ------------------------------------------------------------
    put(ws_mon, "A1", "Monthly net profit by strategy (booked on exit date, 1 MNQ contract)", bold=True, size=14)
    months = pd.period_range(trading_days[0], trading_days[-1], freq="M")
    put(ws_mon, "A3", "Month starts", italic=True)
    put(ws_mon, "A4", "Next month starts", italic=True)
    header_row(ws_mon, 5, ["Strategy"] + [m.strftime("%b %Y") for m in months] + ["Total", "Matches Results tab?"])
    for j, m in enumerate(months, start=2):
        col = get_column_letter(j)
        put(ws_mon, f"{col}3", m.start_time.date(), color=BLUE, fmt=DATE)
        put(ws_mon, f"{col}4", (m + 1).start_time.date(), color=BLUE, fmt=DATE)
    last_m_col = get_column_letter(1 + len(months))
    tot_col, chk_col = get_column_letter(2 + len(months)), get_column_letter(3 + len(months))
    for i, n in enumerate(names):
        r = 6 + i
        put(ws_mon, f"A{r}", f"='Results'!A{R0 + i}", color=GREEN, bold=True)
        for j in range(2, 2 + len(months)):
            col = get_column_letter(j)
            put(ws_mon, f"{col}{r}",
                f'=SUMIFS({TL("P")},{TL("A")},$A{r},{TL("L")},">="&{col}$3,{TL("L")},"<"&{col}$4)', fmt=USD)
        put(ws_mon, f"{tot_col}{r}", f"=SUM(B{r}:{last_m_col}{r})", fmt=USD)
        put(ws_mon, f"{chk_col}{r}", f'=IF(ABS({tot_col}{r}-\'Results\'!I{R0 + i})<0.005,"Yes","NO - check")')
    widths(ws_mon, {"A": 26, **{get_column_letter(c): 13 for c in range(2, 4 + len(months))}})
    bar = BarChart()
    bar.type = "col"
    bar.title = "Net profit by month"
    bar.y_axis.title = "Net $"
    bar.height, bar.width = 10, 22
    last_row = 5 + len(names)
    bar.add_data(Reference(ws_mon, min_col=2, max_col=1 + len(months), min_row=5, max_row=last_row), titles_from_data=True)
    bar.set_categories(Reference(ws_mon, min_col=1, min_row=6, max_row=last_row))
    bar.grouping = "clustered"
    ws_mon.add_chart(bar, f"A{last_row + 3}")

    # ---- Breakdowns ---------------------------------------------------------
    put(ws_brk, "A1", "Breakdowns: where each strategy made or lost money", bold=True, size=14)
    put(ws_brk, "A2", "Caution: slicing a small sample many ways always turns up something that looks good by chance. "
                      "Use these to understand behavior, not to pick a winner.", italic=True, color="C00000")
    r = 4
    for title, col_letter, key in [("By setup type", "B", "setup"), ("By direction", "C", "direction"),
                                   ("By exit reason", "J", "exit_reason")]:
        section(ws_brk, f"A{r}", title, span=6)
        header_row(ws_brk, r + 1, ["Strategy", title.replace("By ", "").capitalize(), "Trades", "Win rate",
                                   "Net $", "Avg per trade $"])
        r += 2
        pairs = trades[["strategy", key]].drop_duplicates()
        pairs = pairs[pairs[key] != "not logged"]
        for _, p in pairs.iterrows():
            put(ws_brk, f"A{r}", p["strategy"], color=BLUE)
            put(ws_brk, f"B{r}", p[key], color=BLUE)
            put(ws_brk, f"C{r}", f"=COUNTIFS({TL('A')},$A{r},{TL(col_letter)},$B{r})", fmt="0")
            put(ws_brk, f"D{r}", f'=IFERROR(COUNTIFS({TL("A")},$A{r},{TL(col_letter)},$B{r},{TL("P")},">0")/C{r},0)',
                fmt=PCT)
            put(ws_brk, f"E{r}", f"=SUMIFS({TL('P')},{TL('A')},$A{r},{TL(col_letter)},$B{r})", fmt=USD)
            put(ws_brk, f"F{r}", f"=IFERROR(E{r}/C{r},0)", fmt=USD)
            r += 1
        r += 1
    widths(ws_brk, {"A": 26, "B": 26, "C": 9, "D": 10, "E": 13, "F": 15})

    # ---- Funnel -------------------------------------------------------------
    put(ws_fun, "A1", "Funnel: why 'Killzone legs + PO3 entry' took so few trades", bold=True, size=14)
    put(ws_fun, "A2", "Each row counts setups that survived every stage above it. Counts come from running the "
                      "strategy (examples/build_analysis_workbook.py -> killzone_po3_strategy.run(funnel=...)).", italic=True)
    header_row(ws_fun, 4, ["Stage", "Setups", "% of previous stage", "What this stage requires"])
    stages = [
        ("killzones_checked", "Killzone windows checked", "Every 09:30 and 20:00 open while flat"),
        ("legs_passing_v1_filters", "Legs passing v1's filters", "Valid/pivotal leg, running against the daily bias"),
        ("zone_reached", "Reached the leg's 2.0 level", "Within 60 minutes (4.5 level on trending days)"),
        ("mss", "Market structure shift", "5-min close back through the last swing"),
        ("entered", "Entered a trade", "Retrace into a gap inside the 1-1.5 STDV zone within 4 hours"),
    ]
    for i, (key, label, why) in enumerate(stages, start=5):
        put(ws_fun, f"A{i}", label)
        put(ws_fun, f"B{i}", funnel.get(key, 0), color=BLUE, fmt="0")
        put(ws_fun, f"C{i}", f'=IFERROR(B{i}/B{i - 1},"")' if i > 5 else "", fmt=PCT)
        put(ws_fun, f"D{i}", why)
    widths(ws_fun, {"A": 30, "B": 10, "C": 18, "D": 62})

    # ---- Method -------------------------------------------------------------
    put(ws_meth, "A1", "Method (procedure)", bold=True, size=14)
    steps = [
        ("1. Collect data", "Downloaded 1-minute candles for Nasdaq-100 futures from Databento (CME Globex). See 'Data'."),
        ("2. Clean data", "Dropped spread contracts, kept the most-traded contract each day, back-adjusted at contract "
                          "rolls (the September roll alone had a ~295-point gap between contracts that would otherwise "
                          "look like a real price move), converted times to New York."),
        ("3. Write each strategy as exact rules", "Turned each method (the trader's manipulation-leg method and the two "
                                                   "Standard Deviation PDFs) into code with no discretion. See 'Strategies'."),
        ("4. Fix everything else", "Same candles, contract size, costs, fill rules and account rules for every strategy. "
                                   "See 'Controlled Variables'."),
        ("5. Run each strategy", "Stepped through every candle in time order. A decision on a candle's close can only "
                                 "trade on the next candle, so no strategy can use information from the future."),
        ("6. Record every trade", "Entry, exit, stop, target, reason, and profit for each trade. See 'Trade Log'."),
        ("7. Measure results", "Win rate, profit factor, net profit, drawdown, monthly profit, stability, and robustness "
                               "-- all as formulas over the Trade Log. See 'Results', 'Monthly', 'Daily Equity'."),
        ("8. Test against the prop-firm evaluation", "Replayed each strategy's daily profit against MyFundedFutures Pro "
                                                      "50K rules, starting a fresh evaluation on every trading day."),
        ("9. Check the work", "124 automated tests (python -m pytest) check the math, including tests that truncate the "
                              "data and confirm earlier decisions don't change (no look-ahead). Sample trades were also "
                              "checked by hand against the raw candles."),
    ]
    header_row(ws_meth, 3, ["Step", "What was done"])
    for i, (a, b) in enumerate(steps, start=4):
        put(ws_meth, f"A{i}", a, bold=True, wrap=True)
        put(ws_meth, f"B{i}", b, wrap=True)
        ws_meth.row_dimensions[i].height = 45
    r = 4 + len(steps) + 1
    section(ws_meth, f"A{r}", "Limitations", span=2)
    for text in [
        "About 15 weeks of data from one period -- too little to prove an edge; more years are needed.",
        "Fills at candle closes; real trading adds slippage, especially on stops.",
        "Several strategies and settings were tried on the same data, which makes the best one look better than it really is.",
        "Some rules in the PDFs can't be coded exactly (e.g. SMT divergence needs a second market's data) and were left out.",
    ]:
        r += 1
        put(ws_meth, f"A{r}", "- " + text, wrap=False)
    widths(ws_meth, {"A": 34, "B": 110})

    # ---- Overview -----------------------------------------------------------
    o = ws_over
    put(o, "A1", "Testing Manipulation-Leg & Standard-Deviation Trading Strategies on Nasdaq-100 Futures",
        bold=True, size=16, color=BLUE, fill=INPUT_FILL)
    put(o, "A2", "(Title is a placeholder -- edit it to your project title.)", italic=True)
    rows = [
        ("Question", "Do rule-based 'manipulation leg' and Standard Deviation (STDV) strategies make money on 1-minute "
                     "Nasdaq-100 futures data -- and could they pass a MyFundedFutures Pro 50K evaluation?"),
        ("Independent variable", "The trading strategy (6 versions -- see 'Strategies')."),
        ("Dependent variables", "Number of trades, win rate, profit factor, net profit, max drawdown, monthly profit, "
                                "and MFF evaluation passes/fails (see 'Results')."),
        ("Controlled variables", "Same data, contract size, commission, fill rules, and account rules for every "
                                 "strategy (see 'Controlled Variables')."),
        ("Materials", "Databento 1-minute futures data (see 'Data'); Python backtesting code (the STDVBOT project); this workbook."),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        put(o, f"A{i}", k, bold=True)
        put(o, f"B{i}", v, wrap=True)
        o.row_dimensions[i].height = 30
    r = 4 + len(rows) + 1
    section(o, f"A{r}", "Headline results (linked from 'Results')", span=7)
    header_row(o, r + 1, ["Strategy", "Trades", "Win rate", "Net profit $", "Max drawdown $",
                          "MFF evals passed", "MFF evals failed"])
    for i, n in enumerate(names):
        rr = r + 2 + i
        src = R0 + i
        put(o, f"A{rr}", f"='Results'!A{src}", color=GREEN, bold=True)
        put(o, f"B{rr}", f"='Results'!B{src}", color=GREEN, fmt="0")
        put(o, f"C{rr}", f"='Results'!E{src}", color=GREEN, fmt=PCT)
        put(o, f"D{rr}", f"='Results'!I{src}", color=GREEN, fmt=USD)
        put(o, f"E{rr}", f"='Results'!M{src}", color=GREEN, fmt=USD)
        put(o, f"F{rr}", f"='Results'!F{mff_rows[n]}", color=GREEN, fmt="0")
        put(o, f"G{rr}", f"='Results'!G{mff_rows[n]}", color=GREEN, fmt="0")
    r = r + 2 + len(names) + 1
    section(o, f"A{r}", "Conclusion so far", span=7)
    for text in [
        "No strategy showed a reliable edge in this sample.",
        "v1/v2 and the hybrid traded too rarely to judge.",
        "PO3 and IPDA made most of their profit in June-July and were roughly flat afterward.",
        "PO3 alone failed the MFF evaluation far more often than it passed across different start dates.",
        "Next step: test on data from other years before choosing or tuning a strategy.",
    ]:
        r += 1
        put(o, f"A{r}", "- " + text)
    r += 2
    section(o, f"A{r}", "How to read this workbook", span=7)
    guide = [
        ("Method", "The procedure, step by step, plus limitations"),
        ("Data", "Every 1-minute candle used (materials)"),
        ("Controlled Variables", "Settings held the same for every strategy (edit the yellow cells to test changes)"),
        ("Strategies", "The independent variable: each strategy's exact rules"),
        ("Trade Log", "Every trade from every strategy"),
        ("Results", "Dependent variables: performance, stability, MFF evaluation"),
        ("Monthly / Daily Equity", "Profit over time, with charts"),
        ("Breakdowns", "Results by setup type, direction, exit reason"),
        ("Funnel", "Why the hybrid strategy took so few trades"),
    ]
    for tab, desc in guide:
        r += 1
        put(o, f"A{r}", tab, bold=True)
        put(o, f"B{r}", desc)
    r += 2
    section(o, f"A{r}", "Color key", span=7)
    for label, color, desc in [("Blue text", BLUE, "Recorded data or an input setting"),
                               ("Black text", BLACK, "Calculated by formula"),
                               ("Green text", GREEN, "Linked from another tab"),
                               ("Yellow fill", BLACK, "Settings worth changing to test 'what if'")]:
        r += 1
        put(o, f"A{r}", label, color=color, bold=True, fill=INPUT_FILL if label == "Yellow fill" else None)
        put(o, f"B{r}", desc)
    widths(o, {"A": 26, "B": 16, "C": 12, "D": 14, "E": 15, "F": 16, "G": 16})
    o.column_dimensions["B"].width = 16
    for i in range(4, 4 + len(rows)):
        o.merge_cells(f"B{i}:G{i}")

    # openpyxl stores formulas without results; make Excel compute them on open.
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--data", required=True)
    parser.add_argument("--out", default="STDVBOT_Analysis.xlsx")
    args = parser.parse_args()
    df = load(args.data)
    trades, funnel = run_all(df)
    build(df, trades, funnel, args.out, Path(args.data).name)
    print(f"Wrote {args.out}: {len(df):,} candles, {len(trades)} trades across {trades['strategy'].nunique()} strategies")


if __name__ == "__main__":
    main()
